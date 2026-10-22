"""Comprehensive report: compiles inspection, cleaning, statistics, trends, models and
Q&A into one document, previewed in the app and exported as HTML and Word.

The report is built as a list of "blocks" (heading, text, table, image, code...) so the
same content feeds the preview, the HTML file and the Word file.
"""
import base64
import html
import platform
import re
import sys
from datetime import date
from io import BytesIO

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import scipy
import seaborn as sns
import sklearn
import streamlit as st
from sklearn.inspection import permutation_importance

import cleaning
import llm
import ml
import project
import qa
from common import analysis_numeric_cols, categorical_cols, id_like_columns, numeric_cols
from eda import detect_date_columns
from inspection import missing_summary, missingness_vs_others, quality_issues
from stats_analysis import (bh_adjust, categorical_summary, chi_square_test, compare_groups, correlation_pairs,
                            descriptive_table, key_findings, normality_tests, outlier_summary, vif_table)

SECTIONS = ["Executive summary", "Dataset overview", "Data quality and missing data", "Data cleaning",
            "Descriptive statistics", "Distributions and normality", "Correlation analysis",
            "Categorical analysis", "Group comparisons", "Outlier analysis", "Time trends",
            "Predictive modelling", "Questions and answers", "Conclusions and notes",
            "Appendix: methods and reproducibility"]
_XML_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def slug(name: str) -> str:
    return re.sub(r"\W+", "_", name.lower()).strip("_")


# ----------------------------------------------------------------------------
# Document model
# ----------------------------------------------------------------------------
def fmt_value(v, col: str = "") -> str:
    if isinstance(v, (bool, np.bool_)):
        return "Yes" if v else "No"
    if isinstance(v, (int, np.integer)):
        return f"{v:,}"
    if isinstance(v, (float, np.floating)):
        if np.isnan(v):
            return ""
        if np.isinf(v):
            return "inf"
        lc = col.lower()
        if ("p-value" in lc or lc.startswith("p (") or lc in ("p", "shapiro p")) and 0 <= v <= 1:
            return "<0.0001" if v < 1e-4 else f"{v:.4f}"
        if abs(v) >= 1000:
            return f"{v:,.0f}"
        if abs(v) >= 100:
            return f"{v:,.1f}"
        return f"{v:.4g}"
    return "" if v is None else str(v)


def pretty(df: pd.DataFrame, max_rows: int = 40, max_cols: int = 10) -> pd.DataFrame:
    out = df.iloc[:max_rows, :max_cols].copy()
    for c in out.columns:
        out[c] = [fmt_value(v, str(c)) for v in out[c]]
    return out


def png(fig) -> bytes:
    buf = BytesIO()
    fig.savefig(buf, format="png", dpi=130, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


class Doc:
    def __init__(self):
        self.blocks = []

    def h1(self, t): self.blocks.append(("h1", t))
    def h2(self, t): self.blocks.append(("h2", t))
    def h3(self, t): self.blocks.append(("h3", t))
    def p(self, t): self.blocks.append(("p", t))
    def bullets(self, items): self.blocks.append(("bullets", [str(i) for i in items]))
    def code(self, t): self.blocks.append(("code", t))
    def callout(self, t): self.blocks.append(("callout", t))
    def image(self, data, caption=""): self.blocks.append(("image", data, caption))

    def table(self, df: pd.DataFrame, caption: str = "", index: bool = False, max_rows: int = 40, max_cols: int = 10):
        if df is None or df.empty:
            return
        shown = df.reset_index() if index and not isinstance(df.index, pd.RangeIndex) else df
        extra = ""
        if len(shown) > max_rows:
            extra += f" (first {max_rows} of {len(shown)} rows)"
        if shown.shape[1] > max_cols:
            extra += f" (first {max_cols} of {shown.shape[1]} columns)"
        self.blocks.append(("table", pretty(shown, max_rows, max_cols), caption + extra))


def _grid(n, ncols=3, height=3.0):
    rows = max(1, (n + ncols - 1) // ncols)
    fig, axes = plt.subplots(rows, ncols, figsize=(4.3 * ncols, height * rows))
    axes = np.atleast_1d(axes).ravel()
    for ax in axes[n:]:
        ax.axis("off")
    return fig, axes


# ----------------------------------------------------------------------------
# Facts shared by the executive summary and the AI drafting step
# ----------------------------------------------------------------------------
def collect_facts(c: dict) -> list:
    df, raw = c["df"], c["raw"]
    facts = [f"Dataset analysed: {c['label']} with {len(df):,} rows and {df.shape[1]} columns."]
    miss = int(raw.isna().sum().sum())
    facts.append(f"The original data had {miss:,} missing cells ({miss / max(raw.size, 1) * 100:.2f}%) "
                 f"and {int(raw.duplicated().sum())} duplicate rows.")
    if c["clean_log"]:
        facts.append(f"Cleaning applied {len(c['clean_log'])} actions; the cleaned data has "
                     f"{int(df.isna().sum().sum())} missing cells.")
    try:
        facts += key_findings(df)[:6]
    except Exception:
        pass
    ms = c.get("ml")
    if ms:
        facts.append(ml.explain_choice(ms["metrics"], ms["problem"]).replace("**", "")
                     + f" Target: {ms['target']} ({ms['problem']}).")
    for r in c.get("qa", [])[:6]:
        if r.get("answer"):
            facts.append(f"Question: {r['question']} Answer: {r['answer'][:300]}")
    return facts


# ----------------------------------------------------------------------------
# Section builders
# ----------------------------------------------------------------------------
def sec_exec(d, c):
    d.h1("Executive Summary")
    text = (st.session_state.get("report_exec") or "").strip()
    if text:
        for para in [p.strip() for p in text.split("\n\n") if p.strip()]:
            d.p(para)
    d.h3("Key facts")
    d.bullets(collect_facts(c))


def sec_overview(d, c):
    df = c["df"]
    d.h1("Dataset Overview")
    dts = detect_date_columns(df)
    d.table(pd.DataFrame({"Item": ["Dataset analysed", "Rows", "Columns", "Numeric columns", "Categorical columns",
                                   "Date columns", "Missing cells", "Duplicate rows"],
                          "Value": [c["label"], f"{len(df):,}", df.shape[1], len(numeric_cols(df)),
                                    len(categorical_cols(df)), len(dts), f"{int(df.isna().sum().sum()):,}",
                                    int(df.duplicated().sum())]}), "Dataset at a glance")
    info = pd.DataFrame({"Column": df.columns, "Type": df.dtypes.astype(str).values,
                         "Non-null": df.notna().sum().values, "Unique": df.nunique().values,
                         "Example": [str(df[x].dropna().iloc[0])[:28] if df[x].notna().any() else "" for x in df.columns]})
    d.table(info, "Column inventory", max_rows=60)
    d.table(df.head(8), "First rows of the data")


def sec_quality(d, c):
    raw = c["raw"]
    d.h1("Data Quality and Missing Data")
    d.p("This section inspects the ORIGINAL data, before any cleaning.")
    ms = missing_summary(raw)
    with_gaps = ms[ms["missing"] > 0].sort_values("% missing", ascending=False)
    total = int(raw.isna().sum().sum())
    d.p(f"{total:,} of {raw.size:,} cells are missing ({total / max(raw.size, 1) * 100:.2f}%); "
        f"{int(raw.isna().any(axis=1).sum()):,} rows contain at least one gap.")
    if with_gaps.empty:
        d.p("No explicit missing values were found.")
    else:
        d.table(with_gaps[["dtype", "missing", "% missing", "severity", "suggested handling"]],
                "Columns with missing values", index=True)
        top = with_gaps.head(15)
        fig, ax = plt.subplots(figsize=(6.5, max(2.2, 0.35 * len(top))))
        sns.barplot(x=top["% missing"], y=top.index, ax=ax, color="indianred")
        ax.set_xlabel("% missing")
        d.image(png(fig), "Percentage of missing values per column")
        lines = []
        for col in with_gaps.index[:3]:
            res = missingness_vs_others(raw, col)
            if res.empty:
                continue
            rel = res[res["p-value"] < 0.05]["compared column"].tolist()
            lines.append(f"{col}: " + (f"missingness is related to {', '.join(rel)} (Missing At Random); "
                                       "imputation is preferable to deleting rows." if rel else
                                       "no relationship with other numeric columns, consistent with Missing "
                                       "Completely At Random; deletion or simple imputation is low-risk."))
        if lines:
            d.h3("Is the missingness random?")
            d.bullets(lines)
    issues = quality_issues(raw)
    if not issues.empty:
        d.table(issues, "Other data-quality findings", max_rows=30)


def sec_cleaning(d, c):
    d.h1("Data Cleaning")
    if not c["clean_log"]:
        d.p("No cleaning was applied; the analysis uses the original data.")
        return
    before = st.session_state.get("clean_before") or {}
    after = {"rows": len(c["cleaned"]), "cols": c["cleaned"].shape[1],
             "missing": int(c["cleaned"].isna().sum().sum()), "dups": int(c["cleaned"].duplicated().sum())}
    d.table(pd.DataFrame({"Measure": ["Rows", "Columns", "Missing cells", "Duplicate rows"],
                          "Before": [before.get(k, "") for k in ("rows", "cols", "missing", "dups")],
                          "After": [after[k] for k in ("rows", "cols", "missing", "dups")]}), "Before and after cleaning")
    d.table(pd.DataFrame({"Step": range(1, len(c["clean_log"]) + 1), "Action taken": [x[0] for x in c["clean_log"]]}),
            "Cleaning log (every decision applied)", max_rows=60)
    d.p("The exact Python code for these steps is in the appendix.")


def sec_descriptive(d, c):
    df, nums = c["df"], c["nums"]
    d.h1("Descriptive Statistics")
    if nums:
        t = descriptive_table(df, nums[:25])
        d.table(t[["count", "missing", "mean", "median", "std", "min", "max", "skewness", "kurtosis (excess)", "CV %"]],
                "Summary of numeric columns", index=True, max_rows=30, max_cols=10)
        d.table(t[["Q1", "Q3", "IQR", "range", "SEM", "95% CI low", "95% CI high"]],
                "Spread and 95% confidence interval of the mean", index=True, max_rows=30)
    cs = categorical_summary(df, c["cats"])
    if not cs.empty:
        d.table(cs, "Categorical columns", index=True)


def sec_distribution(d, c):
    df, nums = c["df"], c["nums"]
    d.h1("Distributions and Normality")
    if not nums:
        d.p("No numeric columns.")
        return
    rows = []
    for col in nums[:25]:
        t = normality_tests(df[col])
        rows.append({"Column": col, "Skewness": df[col].skew(), "Kurtosis": df[col].kurt(),
                     "Shapiro p": t.iloc[0, 2] if not t.empty else np.nan,
                     "Looks normal (alpha 0.05)": bool(t.iloc[0, 3]) if not t.empty else False})
    tbl = pd.DataFrame(rows)
    d.table(tbl, "Normality assessment (Shapiro-Wilk)")
    skewed = tbl[tbl["Skewness"].abs() > 1]["Column"].tolist()
    if skewed:
        d.p(f"Strongly skewed (|skew| > 1): {', '.join(skewed)}. Use medians and rank-based tests for these, "
            "and consider a log transform before modelling.")
    cols = nums[:c["max"]]
    fig, axes = _grid(len(cols))
    for ax, col in zip(axes, cols):
        sns.histplot(df[col].dropna(), kde=True, ax=ax, color="steelblue")
        ax.set_title(col, fontsize=10)
    d.image(png(fig), "Distribution of numeric columns")


def sec_correlation(d, c):
    df = c["df"]
    cols = c["nums"][:15]
    d.h1("Correlation Analysis")
    if len(cols) < 2:
        d.p("At least two numeric columns are needed.")
        return
    fig, ax = plt.subplots(figsize=(7, 5.5))
    sns.heatmap(df[cols].corr(), annot=len(cols) <= 10, fmt=".2f", cmap="coolwarm", vmin=-1, vmax=1, ax=ax)
    d.image(png(fig), "Pearson correlation matrix")
    pairs = correlation_pairs(df, cols)
    if not pairs.empty:
        d.table(pairs[["column A", "column B", "n", "r", "p-value", "p (BH-adjusted)", "strength"]].head(12),
                "Strongest relationships (BH-adjusted p-values control for testing many pairs)")
        top = pairs.iloc[0]
        d.p(f"The strongest relationship is between {top['column A']} and {top['column B']} "
            f"(r = {top['r']:.2f}, {top['strength']}, {top['direction']}).")
    v = vif_table(df, cols)
    if not v.empty:
        d.table(v, "Multicollinearity (VIF above 5 suggests overlap, above 10 is severe)", index=True)


def sec_categorical(d, c):
    df, cats = c["df"], c["cats"]
    d.h1("Categorical Analysis")
    if not cats:
        d.p("No categorical columns with 30 or fewer categories.")
        return
    shown = cats[:min(len(cats), 6)]
    fig, axes = _grid(len(shown))
    for ax, col in zip(axes, shown):
        vc = df[col].value_counts().head(8)
        ax.barh(vc.index.astype(str)[::-1], vc.values[::-1], color="steelblue")
        ax.set_title(col, fontsize=10)
    d.image(png(fig), "Most frequent categories")
    rows = []
    for i, a in enumerate(shown):
        for b in shown[i + 1:]:
            try:
                r = chi_square_test(df[[a, b]].dropna(), a, b)
                rows.append({"Variable A": a, "Variable B": b, "Chi-square": r["chi2"], "p-value": r["p"],
                             "Cramer's V": r["cramers_v"], "Associated (p<0.05)": r["p"] < 0.05})
            except Exception:
                continue
    if rows:
        d.table(pd.DataFrame(rows).sort_values("p-value"), "Chi-square tests of independence")


def sec_groups(d, c):
    df = c["df"]
    nums = c["nums"][:6]
    grp = [g for g in c["cats"] if 2 <= df[g].nunique() <= 10][:4]
    d.h1("Group Comparisons")
    if not nums or not grp:
        d.p("Needs numeric columns and categorical columns with 2 to 10 groups.")
        return
    rows = []
    for g in grp:
        for n in nums:
            r = compare_groups(df, n, g)
            if r is None:
                continue
            normal = all(p > 0.05 for p in r["shapiro"].values() if not np.isnan(p))
            big = min(len(v) for v in r["groups"].values()) >= 30
            t = r["tests"].iloc[0 if (normal or big) else 1]
            rows.append({"Outcome": n, "Grouped by": g, "Groups": r["k"], "Test": t["Test"],
                         "Statistic": t["Statistic"], "p-value": t["p-value"], "Effect size": t["Effect size"]})
    if not rows:
        d.p("No valid comparisons could be made.")
        return
    tbl = pd.DataFrame(rows)
    tbl["p (BH-adjusted)"] = bh_adjust(tbl["p-value"].values)
    tbl["Significant"] = tbl["p (BH-adjusted)"] < 0.05
    tbl = tbl.sort_values("p (BH-adjusted)")
    d.p(f"{int(tbl['Significant'].sum())} of {len(tbl)} comparisons are significant after multiple-testing "
        "adjustment. Parametric tests are used when data are roughly normal or every group has 30+ rows; "
        "otherwise rank-based tests.")
    d.table(tbl[["Outcome", "Grouped by", "Test", "p-value", "p (BH-adjusted)", "Effect size", "Significant"]],
            "Group comparison results", max_rows=25)
    sig = tbl[tbl["Significant"]].head(3)
    if not sig.empty:
        fig, axes = _grid(len(sig), ncols=min(3, len(sig)), height=3.4)
        for ax, (_, r) in zip(axes, sig.iterrows()):
            data = df[[r["Outcome"], r["Grouped by"]]].dropna().astype({r["Grouped by"]: str})
            sns.boxplot(data=data, x=r["Grouped by"], y=r["Outcome"], ax=ax)
            ax.tick_params(axis="x", rotation=30, labelsize=8)
        d.image(png(fig), "Most significant differences")


def sec_outliers(d, c):
    df, nums = c["df"], c["nums"]
    d.h1("Outlier Analysis")
    if not nums:
        d.p("No numeric columns.")
        return
    d.table(outlier_summary(df, nums[:25]), "Outliers by three methods (IQR rule, z-score, modified z-score)", index=True)
    cols = nums[:c["max"]]
    fig, axes = _grid(len(cols))
    for ax, col in zip(axes, cols):
        sns.boxplot(x=df[col].dropna(), ax=ax, color="lightcoral")
        ax.set_title(col, fontsize=10)
    d.image(png(fig), "Box plots")


def sec_trends(d, c):
    df = c["df"]
    d.h1("Time Trends")
    dcols, nums = detect_date_columns(df), c["nums"]
    if not dcols or not nums:
        d.p("No date column was found, so no time analysis was produced.")
        return
    dc = dcols[0]
    preferred = ("sales", "revenue", "amount", "total", "price", "value", "profit", "cost", "income", "score", "count")
    vc = next((n for p in preferred for n in nums if p in n.lower()), nums[0])
    ts = df[[dc, vc]].copy()
    ts[dc] = pd.to_datetime(ts[dc], errors="coerce")
    ts = ts.dropna(subset=[dc])
    if ts.empty:
        d.p("The date column could not be parsed.")
        return
    monthly = ts.groupby(ts[dc].dt.to_period("M").dt.to_timestamp())[vc].sum().sort_index()
    d.p(f"Total {vc} by month, using {dc} ({ts[dc].min().date()} to {ts[dc].max().date()}).")
    fig, ax = plt.subplots(figsize=(8, 3.5))
    ax.plot(monthly.index, monthly.values, marker="o", markersize=3, label="Monthly total")
    if len(monthly) >= 3:
        ax.plot(monthly.index, monthly.rolling(3).mean().values, color="crimson", linewidth=2, label="3-month average")
    ax.legend()
    ax.grid(alpha=0.3)
    d.image(png(fig), f"Monthly {vc}")
    seas = ts.groupby(ts[dc].dt.month)[vc].sum().reindex(range(1, 13))
    seas.index = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    fig, ax = plt.subplots(figsize=(7, 3))
    ax.bar(seas.index, seas.fillna(0).values, color="steelblue")
    d.image(png(fig), f"Seasonality: {vc} by calendar month")
    d.bullets([f"Highest month: {monthly.idxmax():%b %Y} ({monthly.max():,.0f}).",
               f"Lowest month: {monthly.idxmin():%b %Y} ({monthly.min():,.0f}).",
               f"Strongest calendar month overall: {seas.idxmax()}."])
    yearly = ts.groupby(ts[dc].dt.year)[vc].sum()
    if len(yearly) > 1:
        yt = pd.DataFrame({"Year": yearly.index.astype(str), f"Total {vc}": yearly.values,
                           "Change vs previous year %": (yearly.pct_change() * 100).values})
        d.table(yt, "Year-over-year totals")


def _ensure_perm(res, name):
    if name in res["perm"] or name not in res.get("models", {}):
        return
    Xs = res["X_test"].sample(min(1000, len(res["X_test"])), random_state=42)
    pi = permutation_importance(res["models"][name], Xs, res["y_test"].loc[Xs.index],
                                scoring="r2" if res["problem"] == "Regression" else "f1_weighted",
                                n_repeats=5, random_state=42, n_jobs=1)
    res["perm"][name] = pd.DataFrame({"feature": Xs.columns, "importance": pi.importances_mean,
                                      "std": pi.importances_std}).sort_values("importance", ascending=False)


def sec_ml(d, c):
    s = c.get("ml")
    d.h1("Predictive Modelling")
    if not s:
        d.p("No predictive models were trained in this session.")
        return
    prim = ml.primary_metric(s["problem"])
    m = s["metrics"].sort_values("CV score (mean)", ascending=False)
    d.p(f"Problem type: {s['problem']}. Target: {s['target']}. Features used: {', '.join(map(str, s['kept']))}. "
        f"{len(m)} models were compared using cross-validated {prim} on the training data and then scored on a "
        "held-out test set.")
    show = m.copy()
    show.insert(0, "Rank", range(1, len(show) + 1))
    d.table(show, "Model comparison (sorted by cross-validated score)", index=True, max_rows=30, max_cols=12)
    d.h3("Which model is best, and why")
    d.p(ml.explain_choice(m, s["problem"]).replace("**", ""))
    fig, ax = plt.subplots(figsize=(7, max(2.5, 0.35 * len(m))))
    ax.barh(m.index[::-1], m["CV score (mean)"].values[::-1], xerr=m["CV score (std)"].values[::-1], color="seagreen")
    ax.set_xlabel(f"Cross-validated {prim}")
    d.image(png(fig), "Cross-validated score by model (error bars: standard deviation)")
    best = m.index[0]
    imp = s["perm"].get(best)
    if imp is not None and not imp.empty:
        top = imp.head(10)
        fig, ax = plt.subplots(figsize=(6.5, 3.5))
        sns.barplot(data=top, x="importance", y="feature", ax=ax, color="seagreen")
        d.image(png(fig), f"Feature importance for {best} (permutation importance)")
    if s.get("tuned_params"):
        d.p("Best hyperparameters from tuning: " + ", ".join(f"{k.replace('model__', '')} = {v}"
                                                              for k, v in s["tuned_params"].items()))
    d.callout("Scores come from one train/test split. Confirm on new data before relying on the model for decisions.")


def sec_qa(d, c):
    d.h1("Questions and Answers")
    items = c.get("qa", [])
    if not items:
        d.p("No questions were answered in this session.")
        return
    for i, r in enumerate(items, 1):
        d.h2(f"Q{i}. {r['question']}")
        if r.get("error"):
            d.callout(f"This question could not be fully answered: {r['error']}")
        if r.get("answer"):
            d.p(r["answer"])
        if r.get("approach"):
            d.h3("Procedure")
            d.bullets([f"Step {n}: {s}" for n, s in enumerate(r["approach"], 1)])
        if r.get("libraries"):
            d.table(pd.DataFrame(r["libraries"], columns=["Library", "Why it is needed"]), "Libraries used")
        if r.get("code"):
            d.h3("Python code")
            d.code(r["code"])
        if r.get("stdout"):
            d.h3("Printed output")
            d.code(r["stdout"])
        if isinstance(r.get("result"), pd.DataFrame):
            d.table(r["result"], "Result", index=True, max_rows=25)
        elif r.get("result") is not None:
            d.p(f"Result: {r['result']}")
        for img in r.get("images", []):
            d.image(img, "Figure produced by the analysis")


def sec_conclusions(d, c):
    d.h1("Conclusions and Notes")
    text = (st.session_state.get(f"note_{slug('Conclusions and notes')}") or "").strip()
    if text:
        for para in [p.strip() for p in text.split("\n\n") if p.strip()]:
            d.p(para)
    else:
        d.p("No analyst conclusions were entered. Add them in the report settings before generating.")


def sec_appendix(d, c):
    d.h1("Appendix: Methods and Reproducibility")
    d.table(pd.DataFrame([
        ("Missing-data inspection", "pandas isna, Welch t-test (scipy.stats)", "Find gaps and test whether they are random"),
        ("Cleaning and imputation", "pandas, scikit-learn KNNImputer, StandardScaler", "Fix types, fill or drop missing values, treat outliers"),
        ("Descriptive statistics", "pandas, numpy, scipy.stats", "Centre, spread, shape, confidence intervals"),
        ("Normality", "scipy.stats: Shapiro-Wilk, D'Agostino, Jarque-Bera, Anderson-Darling", "Choose parametric or rank-based methods"),
        ("Correlation", "scipy.stats: Pearson, Spearman, Kendall; Benjamini-Hochberg adjustment", "Measure and test relationships"),
        ("Multicollinearity", "scikit-learn LinearRegression (VIF)", "Detect overlapping predictors"),
        ("Association and group tests", "scipy.stats: chi-square, t-test, Mann-Whitney, ANOVA, Kruskal-Wallis, Tukey HSD", "Test differences and dependence"),
        ("Modelling", "scikit-learn pipelines, cross_validate, RandomizedSearchCV; XGBoost (optional)", "Train, compare, tune and select models"),
        ("Question answering", "An AI model (Google Gemini free tier or Claude) writes pandas code, executed locally", "Answer questions with a visible procedure"),
    ], columns=["Task", "Methods and libraries", "Purpose"]), "Methods used", max_rows=20)
    d.p(f"Software: Python {platform.python_version()}, pandas {pd.__version__}, numpy {np.__version__}, "
        f"scipy {scipy.__version__}, scikit-learn {sklearn.__version__}.")
    if c["clean_log"]:
        d.h3("Cleaning script (reproduces the cleaned dataset)")
        d.code(cleaning.build_script(c["clean_log"]))
    d.h3("Limitations")
    d.bullets(["Statistical significance is not proof of causation.",
               "Imputed values are estimates; conclusions on heavily imputed columns are less certain.",
               "Model performance on a single hold-out split can differ on new data.",
               "Answers generated by an AI model should be verified against the code shown."])


BUILDERS = {"Executive summary": sec_exec, "Dataset overview": sec_overview,
            "Data quality and missing data": sec_quality, "Data cleaning": sec_cleaning,
            "Descriptive statistics": sec_descriptive, "Distributions and normality": sec_distribution,
            "Correlation analysis": sec_correlation, "Categorical analysis": sec_categorical,
            "Group comparisons": sec_groups, "Outlier analysis": sec_outliers, "Time trends": sec_trends,
            "Predictive modelling": sec_ml, "Questions and answers": sec_qa,
            "Conclusions and notes": sec_conclusions, "Appendix: methods and reproducibility": sec_appendix}


def build_blocks(c: dict, sections: list, progress=None) -> list:
    doc = Doc()
    chosen = [s for s in SECTIONS if s in sections]
    for i, name in enumerate(chosen):
        if progress:
            progress(i / max(len(chosen), 1), name)
        before = len(doc.blocks)
        try:
            BUILDERS[name](doc, c)
        except Exception as e:  # noqa: BLE001 - one failing section must not lose the report
            del doc.blocks[before:]
            doc.h1(name)
            doc.callout(f"This section could not be generated: {type(e).__name__}: {e}")
        note = (st.session_state.get(f"note_{slug(name)}") or "").strip()
        if note and name not in ("Conclusions and notes",):
            doc.callout(f"Analyst note: {note}")
    return doc.blocks


# ----------------------------------------------------------------------------
# Renderers
# ----------------------------------------------------------------------------
CSS = """body{font-family:Segoe UI,Arial,sans-serif;max-width:980px;margin:30px auto;padding:0 20px;color:#222;line-height:1.5}
h1{border-bottom:2px solid #2f6f9f;padding-bottom:4px;margin-top:42px;color:#1f4e79}h2{color:#2f6f9f;margin-top:30px}
h3{color:#444}table{border-collapse:collapse;font-size:12px;margin:8px 0 4px;width:100%}
th{background:#2f6f9f;color:#fff;text-align:left;padding:5px 7px}td{border-bottom:1px solid #ddd;padding:4px 7px}
tr:nth-child(even) td{background:#f5f8fb}.cap{font-size:12px;color:#666;font-style:italic;margin:2px 0 16px}
pre{background:#f4f4f4;border:1px solid #ddd;padding:10px;overflow-x:auto;font-size:12px}img{max-width:100%}
.callout{background:#fff8e1;border-left:4px solid #f0b400;padding:8px 12px;margin:10px 0}
.title{font-size:32px;color:#1f4e79;margin-bottom:0}.meta{color:#666;margin-top:4px}nav{background:#f5f8fb;padding:10px 20px;border-radius:6px}"""


def to_html(blocks: list, meta: dict) -> bytes:
    e = html.escape
    toc = [b[1] for b in blocks if b[0] == "h1"]
    out = [f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>{e(meta['title'])}</title><style>{CSS}</style></head><body>",
           f"<div class='title'>{e(meta['title'])}</div>",
           f"<div class='meta'>{e(meta['author'])}{' | ' if meta['author'] else ''}{meta['date']} | Dataset: {e(meta['dataset'])}</div>",
           "<nav><b>Contents</b><ol>" + "".join(f"<li><a href='#s{i}'>{e(t)}</a></li>" for i, t in enumerate(toc)) + "</ol></nav>"]
    n = 0
    for b in blocks:
        k = b[0]
        if k == "h1":
            out.append(f"<h1 id='s{n}'>{e(b[1])}</h1>")
            n += 1
        elif k in ("h2", "h3"):
            out.append(f"<{k}>{e(b[1])}</{k}>")
        elif k == "p":
            out.append(f"<p>{e(b[1])}</p>")
        elif k == "bullets":
            out.append("<ul>" + "".join(f"<li>{e(i)}</li>" for i in b[1]) + "</ul>")
        elif k == "callout":
            out.append(f"<div class='callout'>{e(b[1])}</div>")
        elif k == "code":
            out.append(f"<pre>{e(b[1])}</pre>")
        elif k == "table":
            out.append(b[1].to_html(index=False, border=0, escape=True))
            out.append(f"<div class='cap'>{e(b[2])}</div>" if b[2] else "")
        elif k == "image":
            out.append(f"<p><img src='data:image/png;base64,{base64.b64encode(b[1]).decode()}'></p>")
            out.append(f"<div class='cap'>{e(b[2])}</div>" if b[2] else "")
    out.append("</body></html>")
    return "\n".join(out).encode("utf-8")


def _clean_xml(t) -> str:
    return _XML_BAD.sub("", str(t))


def to_docx(blocks: list, meta: dict) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.8)
    doc.add_heading(_clean_xml(meta["title"]), 0)
    doc.add_paragraph(_clean_xml(f"{meta['author']}{' | ' if meta['author'] else ''}{meta['date']} | Dataset: {meta['dataset']}"))
    doc.add_paragraph("Contents", style="Heading 2")
    for t in [b[1] for b in blocks if b[0] == "h1"]:
        doc.add_paragraph(_clean_xml(t), style="List Number")
    doc.add_page_break()
    for b in blocks:
        k = b[0]
        if k == "h1":
            doc.add_heading(_clean_xml(b[1]), 1)
        elif k == "h2":
            doc.add_heading(_clean_xml(b[1]), 2)
        elif k == "h3":
            doc.add_heading(_clean_xml(b[1]), 3)
        elif k == "p":
            doc.add_paragraph(_clean_xml(b[1]))
        elif k == "bullets":
            for item in b[1]:
                doc.add_paragraph(_clean_xml(item), style="List Bullet")
        elif k == "callout":
            para = doc.add_paragraph()
            run = para.add_run(_clean_xml(b[1]))
            run.italic = True
        elif k == "code":
            para = doc.add_paragraph()
            run = para.add_run(_clean_xml(b[1]))
            run.font.name = "Consolas"
            run.font.size = Pt(7.5)
        elif k == "table":
            df = b[1]
            table = doc.add_table(rows=1, cols=len(df.columns))
            table.style = "Table Grid"
            for j, col in enumerate(df.columns):
                cell = table.rows[0].cells[j]
                cell.text = _clean_xml(col)
                for r in cell.paragraphs[0].runs:
                    r.font.bold, r.font.size = True, Pt(8)
            for _, row in df.iterrows():
                cells = table.add_row().cells
                for j, v in enumerate(row):
                    cells[j].text = _clean_xml(v)
                    for r in cells[j].paragraphs[0].runs:
                        r.font.size = Pt(8)
            if b[2]:
                cap = doc.add_paragraph()
                r = cap.add_run(_clean_xml(b[2]))
                r.italic, r.font.size = True, Pt(8)
        elif k == "image":
            doc.add_picture(BytesIO(b[1]), width=Inches(6.3))
            if b[2]:
                cap = doc.add_paragraph()
                r = cap.add_run(_clean_xml(b[2]))
                r.italic, r.font.size = True, Pt(8)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def render_blocks(blocks: list):
    for b in blocks:
        k = b[0]
        if k == "h1":
            st.header(b[1])
        elif k == "h2":
            st.subheader(b[1])
        elif k == "h3":
            st.markdown(f"**{b[1]}**")
        elif k == "p":
            st.markdown(b[1].replace("$", "\\$"))
        elif k == "bullets":
            st.markdown("\n".join(f"- {i}".replace("$", "\\$") for i in b[1]))
        elif k == "callout":
            st.info(b[1])
        elif k == "code":
            st.code(b[1], language="python")
        elif k == "table":
            st.dataframe(b[1], hide_index=True)
            if b[2]:
                st.caption(b[2])
        elif k == "image":
            st.image(b[1], caption=b[2] or None)


# ----------------------------------------------------------------------------
# UI
# ----------------------------------------------------------------------------
def _init(key, default):
    if key not in st.session_state:
        st.session_state[key] = default


def _only_valid(key, options):
    if key in st.session_state:
        st.session_state[key] = [v for v in st.session_state[key] if v in options]


def draft_exec_summary(facts: list) -> str:
    """Ask the configured AI model to write the executive summary from the facts only. Returns a status message."""
    creds = llm.get_credentials()
    if not creds.key:
        return "Enter an API key on the Questions & Answers page first (a free Gemini key works)."
    try:
        text = qa.ask_model(
            llm.make_client(creds), creds.model,
            "You write executive summaries for non-technical readers. Use ONLY the facts supplied; never invent "
            "numbers or findings.",
            [{"role": "user", "content": "Write a 150 to 220 word executive summary, then 3 to 5 bullet-point "
                                         "recommendations.\n\nFacts:\n" + "\n".join(f"- {f}" for f in facts)}],
            max_tokens=800)
        st.session_state["report_exec"] = text
        return "Draft written. Edit it freely before generating the report."
    except llm.LLMError as e:
        return f"Could not draft the summary: {e}"
    except Exception as e:  # noqa: BLE001
        return f"Could not draft the summary: {e}"


def _draft_exec(facts: list):
    """Button callback for the Report page."""
    st.session_state["report_exec_msg"] = draft_exec_summary(facts)


def ensure_defaults(df: pd.DataFrame):
    """Initialise report settings (also used when the one-click analysis runs before the Report page is opened)."""
    nums = analysis_numeric_cols(df) or numeric_cols(df)
    ids = set(id_like_columns(df))
    cats = [c for c in categorical_cols(df, 30) if c not in ids]
    _init("report_title", "Data Analysis Report")
    _init("report_author", "")
    _init("report_sections", list(SECTIONS))
    _init("report_max", 8)
    _init("report_nums", nums[:8])
    _init("report_cats", cats[:4])
    _only_valid("report_sections", SECTIONS)
    _only_valid("report_nums", numeric_cols(df))
    _only_valid("report_cats", categorical_cols(df, 30))
    return nums, cats


def generate_report(df_raw, df, label, progress=None) -> list:
    """Build the report from everything in session state and store the preview, HTML and Word versions.
    Returns a list of warning strings."""
    ensure_defaults(df)
    ctx = _context(df_raw, df, label)
    blocks = build_blocks(ctx, st.session_state["report_sections"], progress)
    meta = {"title": st.session_state["report_title"] or "Data Analysis Report",
            "author": st.session_state["report_author"], "date": date.today().isoformat(), "dataset": label}
    warns = []
    st.session_state["report_blocks"] = blocks
    st.session_state["report_html"] = to_html(blocks, meta)
    try:
        st.session_state["report_docx"] = to_docx(blocks, meta)
    except Exception as e:  # noqa: BLE001
        st.session_state["report_docx"] = None
        warns.append(f"Word export failed ({e}); the HTML report is still available.")
    st.session_state["report_name"] = slug(meta["title"]) or "report"
    return warns


def _context(df_raw, df, label) -> dict:
    nums = [x for x in st.session_state.get("report_nums", numeric_cols(df)) if x in df.columns]
    cats = [x for x in st.session_state.get("report_cats", categorical_cols(df, 30)) if x in df.columns]
    res = st.session_state.get("ml_result")
    summary = None
    if res:
        best = res["metrics"]["CV score (mean)"].idxmax()
        try:
            _ensure_perm(res, best)
        except Exception:
            pass
        summary = project.ml_summary_live(res, st.session_state.get("ml_tuned_params"))
    elif st.session_state.get("ml_saved"):
        summary = st.session_state["ml_saved"]
    results = st.session_state.get("qa_results", {})
    mine = [r for (lab, shape, _), r in results.items() if lab == label and tuple(shape) == tuple(df.shape)]
    mine = mine or list(results.values())
    cleaned = st.session_state.get("df_clean")
    return {"df": df, "raw": df_raw, "cleaned": cleaned if cleaned is not None else df, "label": label,
            "nums": nums, "cats": cats, "max": int(st.session_state.get("report_max", 8)),
            "clean_log": st.session_state.get("clean_log", []), "ml": summary, "qa": mine}


def render_report(df_raw: pd.DataFrame, df: pd.DataFrame, label: str, dataset_name: str):
    st.subheader("Comprehensive Report")
    st.caption("Compiles everything you have done into one document. Your cleaning, models and answered "
               "questions are kept for the session; use 'Save your analysis' to keep them for later.")

    ensure_defaults(df)
    nums, cats = numeric_cols(df), categorical_cols(df, 30)

    a, b = st.columns(2)
    a.text_input("Report title", key="report_title")
    b.text_input("Author", key="report_author")
    st.multiselect("Sections to include", SECTIONS, key="report_sections")
    with st.expander("Scope: which columns and how many charts"):
        st.multiselect("Numeric columns to analyse", nums, key="report_nums")
        st.multiselect("Categorical columns to analyse", cats, key="report_cats")
        st.slider("Maximum charts per grid", 3, 12, key="report_max")

    ctx = _context(df_raw, df, label)
    with st.expander("Executive summary (optional, editable)"):
        st.text_area("Summary text shown at the top of the report", key="report_exec", height=180)
        st.button("Draft with AI from the findings", on_click=_draft_exec, args=(collect_facts(ctx),))
        if st.session_state.get("report_exec_msg"):
            st.caption(st.session_state["report_exec_msg"])

    with st.expander("Your notes and conclusions (added to the matching sections)"):
        for s in st.session_state["report_sections"]:
            st.text_area(s, key=f"note_{slug(s)}", height=70)

    if st.button("Generate report", type="primary", disabled=not st.session_state["report_sections"]):
        bar = st.progress(0.0, text="Building report...")
        for w in generate_report(df_raw, df, label, lambda f, n: bar.progress(min(f, 1.0), text=f"Building: {n}")):
            st.warning(w)
        bar.empty()
        st.success("Report generated.")

    if st.session_state.get("report_blocks"):
        name = st.session_state["report_name"]
        d1, d2 = st.columns(2)
        d1.download_button("Download report (HTML, opens in any browser; print to PDF)",
                           st.session_state["report_html"], f"{name}.html", "text/html")
        if st.session_state.get("report_docx"):
            d2.download_button("Download report (Word .docx)", st.session_state["report_docx"], f"{name}.docx",
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        with st.expander("Preview the report", expanded=True):
            render_blocks(st.session_state["report_blocks"])

    st.markdown("---")
    st.markdown("#### Save your analysis")
    st.caption("Saves the original and cleaned data, cleaning log, model comparison, answered questions, notes and "
               "report settings in one file. Reload it later from the sidebar. Trained model objects are not stored "
               "(for safety); retrain if you need to predict again.")
    if st.button("Build project file"):
        st.session_state["project_bytes"] = project.save_project(df_raw, st.session_state, dataset_name)
    if st.session_state.get("project_bytes"):
        st.download_button("Download project (.zip)", st.session_state["project_bytes"],
                           f"{slug(dataset_name) or 'analysis'}_project.zip", "application/zip")
